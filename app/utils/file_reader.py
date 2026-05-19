"""统一文件读取模块 - 所有模块共用"""
import os
import json
import logging
import pandas as pd

logger = logging.getLogger("data_analyzer.file_reader")

# 支持的输入格式
SUPPORTED_INPUT_FORMATS = (
    "所有支持的格式 (*.csv *.xlsx *.xls *.json *.tsv *.parquet "
    "*.html *.db *.sqlite *.feather *.md *.pdf *.docx);;"
    "CSV (*.csv);;Excel (*.xlsx *.xls);;JSON (*.json);;TSV (*.tsv);;"
    "Parquet (*.parquet);;HTML (*.html);;SQLite (*.db *.sqlite);;"
    "Feather (*.feather);;Markdown (*.md);;PDF (*.pdf);;Word (*.docx);;所有文件 (*)"
)


def read_file(path, encoding="utf-8", pdf_engine="auto", pdf_password=""):
    """
    统一文件读取入口

    Args:
        path: 文件路径
        encoding: 文本编码
        pdf_engine: PDF引擎选择 ("auto", "pymupdf4llm", "mineru")
        pdf_password: PDF密码

    Returns:
        pandas DataFrame
    """
    ext = os.path.splitext(path)[1].lower()

    if ext == ".csv":
        return pd.read_csv(path, encoding=encoding)
    elif ext in (".xlsx", ".xls"):
        return pd.read_excel(path)
    elif ext == ".json":
        return _read_json(path, encoding)
    elif ext == ".tsv":
        return pd.read_csv(path, sep="\t", encoding=encoding)
    elif ext == ".parquet":
        return pd.read_parquet(path)
    elif ext in (".html", ".htm"):
        dfs = pd.read_html(path)
        return dfs[0] if dfs else pd.DataFrame()
    elif ext in (".db", ".sqlite"):
        return _read_sqlite(path)
    elif ext == ".feather":
        return pd.read_feather(path)
    elif ext == ".md":
        return _read_markdown_table(path, encoding)
    elif ext == ".pdf":
        from app.utils.pdf_parser import read_pdf
        return read_pdf(path, engine=pdf_engine, password=pdf_password)
    elif ext == ".docx":
        return _read_docx(path)
    else:
        raise ValueError(f"不支持的文件格式: {ext}")


def _read_json(path, encoding="utf-8"):
    """读取 JSON 文件"""
    with open(path, "r", encoding=encoding) as f:
        data = json.load(f)
    if isinstance(data, list):
        return pd.json_normalize(data)
    elif isinstance(data, dict):
        for key, val in data.items():
            if isinstance(val, list) and val and isinstance(val[0], dict):
                return pd.json_normalize(val)
        return pd.json_normalize(data)
    return pd.DataFrame([data])


def _read_sqlite(path):
    """读取 SQLite 文件"""
    import sqlite3
    conn = sqlite3.connect(path)
    tables = pd.read_sql("SELECT name FROM sqlite_master WHERE type='table'", conn)
    if tables.empty:
        conn.close()
        raise ValueError("SQLite 文件中没有表")
    table_name = tables.iloc[0, 0]
    df = pd.read_sql(f"SELECT * FROM [{table_name}]", conn)
    conn.close()
    return df


def _read_markdown_table(path, encoding="utf-8"):
    """解析 Markdown 表格"""
    with open(path, "r", encoding=encoding) as f:
        lines = f.readlines()

    table_lines = [l.strip() for l in lines if "|" in l]
    if len(table_lines) < 2:
        raise ValueError("未找到 Markdown 表格")

    headers = [c.strip() for c in table_lines[0].split("|") if c.strip()]
    data_lines = [l for l in table_lines[1:] if not all(c in "-| " for c in l)]
    rows = []
    for line in data_lines:
        cells = [c.strip() for c in line.split("|") if c.strip()]
        if len(cells) == len(headers):
            rows.append(cells)

    return pd.DataFrame(rows, columns=headers)


def _read_docx(path):
    """读取 Word 文档中的表格"""
    from docx import Document

    doc = Document(path)
    all_dfs = []

    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(cells)
        if len(rows) >= 2:
            all_dfs.append(pd.DataFrame(rows[1:], columns=rows[0]))

    if all_dfs:
        return pd.concat(all_dfs, ignore_index=True)

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    if paragraphs:
        return pd.DataFrame({"内容": paragraphs})

    raise ValueError("Word 文档中未找到表格或文本内容")
