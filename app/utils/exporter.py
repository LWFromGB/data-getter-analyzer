"""统一导出模块 - 所有模块共用"""
import os
import logging
import pandas as pd

logger = logging.getLogger("data_analyzer.exporter")

# 支持的输出格式
OUTPUT_FORMATS = {
    "CSV (.csv)": "csv",
    "Excel (.xlsx)": "xlsx",
    "JSON (.json)": "json",
    "TSV (.tsv)": "tsv",
    "Parquet (.parquet)": "parquet",
    "HTML (.html)": "html",
    "SQLite (.db)": "sqlite",
    "Feather (.feather)": "feather",
    "Markdown (.md)": "md",
    "LaTeX (.tex)": "latex",
    "XML (.xml)": "xml",
    "Word (.docx)": "docx",
    "PDF (.pdf)": "pdf",
}

EXT_FILTERS = {
    "csv": "CSV (*.csv)", "xlsx": "Excel (*.xlsx)", "json": "JSON (*.json)",
    "tsv": "TSV (*.tsv)", "parquet": "Parquet (*.parquet)", "html": "HTML (*.html)",
    "sqlite": "SQLite (*.db)", "feather": "Feather (*.feather)",
    "md": "Markdown (*.md)", "latex": "LaTeX (*.tex)", "xml": "XML (*.xml)",
    "docx": "Word (*.docx)", "pdf": "PDF (*.pdf)",
}


def export_dataframe(df, path, fmt, encoding="utf-8-sig", json_orient="records", table_name="data"):
    """
    统一导出 DataFrame 到指定格式

    Args:
        df: pandas DataFrame
        path: 输出文件路径
        fmt: 格式标识 (csv, xlsx, json, etc.)
        encoding: 输出编码
        json_orient: JSON 方向
        table_name: SQLite 表名
    """
    if fmt == "csv":
        df.to_csv(path, index=False, encoding=encoding)
    elif fmt == "xlsx":
        df.to_excel(path, index=False)
    elif fmt == "json":
        df.to_json(path, orient=json_orient, force_ascii=False, indent=2)
    elif fmt == "tsv":
        df.to_csv(path, index=False, sep="\t", encoding=encoding)
    elif fmt == "parquet":
        df.to_parquet(path, index=False)
    elif fmt == "html":
        html = df.to_html(index=False, classes="table", border=1)
        full_html = (
            "<!DOCTYPE html><html><head><meta charset='utf-8'><title>数据表格</title>"
            "<style>body{font-family:sans-serif;margin:20px;}"
            ".table{border-collapse:collapse;width:100%;}"
            ".table th,.table td{border:1px solid #ddd;padding:8px;text-align:left;}"
            ".table th{background:#4a90d9;color:white;}"
            ".table tr:nth-child(even){background:#f9f9f9;}</style>"
            f"</head><body>{html}</body></html>"
        )
        with open(path, "w", encoding="utf-8") as f:
            f.write(full_html)
    elif fmt == "sqlite":
        import sqlite3
        conn = sqlite3.connect(path)
        df.to_sql(table_name, conn, if_exists="replace", index=False)
        conn.close()
    elif fmt == "feather":
        df.to_feather(path)
    elif fmt == "md":
        with open(path, "w", encoding="utf-8") as f:
            f.write(df.to_markdown(index=False))
    elif fmt == "latex":
        with open(path, "w", encoding="utf-8") as f:
            f.write(df.to_latex(index=False))
    elif fmt == "xml":
        df.to_xml(path, index=False)
    elif fmt == "docx":
        export_docx(df, path)
    elif fmt == "pdf":
        export_pdf(df, path)
    else:
        raise ValueError(f"不支持的输出格式: {fmt}")


def export_docx(df, path):
    """导出为 Word 表格"""
    from docx import Document

    doc = Document()
    doc.add_heading("数据表格", level=1)
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Table Grid"

    for i, col in enumerate(df.columns):
        table.rows[0].cells[i].text = str(col)

    for _, row in df.iterrows():
        row_cells = table.add_row().cells
        for i, val in enumerate(row):
            row_cells[i].text = str(val) if pd.notna(val) else ""

    doc.save(path)


def export_pdf(df, path):
    """导出为 PDF 表格"""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    with PdfPages(path) as pdf:
        page_size = 40
        total_pages = (len(df) + page_size - 1) // page_size

        for page in range(total_pages):
            start = page * page_size
            end = min(start + page_size, len(df))
            page_df = df.iloc[start:end]

            fig, ax = plt.subplots(figsize=(11, 8))
            ax.axis("off")
            ax.set_title(f"数据 (第 {page+1}/{total_pages} 页)", fontsize=12, pad=20)
            col_labels = [str(c) for c in page_df.columns]
            cell_text = [[str(v) if pd.notna(v) else "" for v in row] for _, row in page_df.iterrows()]
            table = ax.table(cellText=cell_text, colLabels=col_labels, loc="center", cellLoc="left")
            table.auto_set_font_size(False)
            table.set_fontsize(7)
            table.scale(1.2, 1.3)
            pdf.savefig(fig, bbox_inches="tight")
            plt.close(fig)


def get_file_size_str(path):
    """获取文件大小的可读字符串"""
    size = os.path.getsize(path)
    if size < 1024:
        return f"{size} B"
    elif size < 1024 * 1024:
        return f"{size/1024:.1f} KB"
    else:
        return f"{size/1024/1024:.1f} MB"
